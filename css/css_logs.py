from selenium.webdriver import ActionChains
from selenium.webdriver.common.by import By
from selenium import webdriver
import time
from selenium.webdriver.support.ui import Select
from selenium.webdriver.support.ui import WebDriverWait
from selenium.webdriver.support import expected_conditions as EC
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter
from datetime import datetime
from selenium.common.exceptions import UnexpectedAlertPresentException, NoAlertPresentException

# Excel Logger Class with Formatting
class ExcelLogger:
    def __init__(self, filename=r"C:\Users\ADMIN\Desktop\DAE\CSS_Report101.xlsx"):
        self.filename = filename
        self.wb = openpyxl.Workbook()
        self.sheet = self.wb.active
        self.sheet.title = "Test Logs"

        headers = ["Timestamp", "Status", "Action", "Data"]
        self.sheet.append(headers)

        # Format header
        for col_num, header in enumerate(headers, 1):
            cell = self.sheet.cell(row=1, column=col_num)
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill(start_color="4F81BD", end_color="4F81BD", fill_type="solid")
            cell.alignment = Alignment(horizontal="center", vertical="center")

        # Set column widths
        col_widths = [20, 10, 40, 60]
        for i, width in enumerate(col_widths, 1):
            self.sheet.column_dimensions[get_column_letter(i)].width = width

        self.wb.save(self.filename)

    def log(self, status, action, data=""):
        now = datetime.now().strftime("%Y-%m-%d %H:%M:%S")
        self.sheet.append([now, status, action, data])

        last_row = self.sheet.max_row
        status_cell = self.sheet.cell(row=last_row, column=2)

        # Status color coding
        if status.upper() == "PASS":
            status_cell.fill = PatternFill(start_color="C6EFCE", end_color="C6EFCE", fill_type="solid")
            status_cell.font = Font(color="006100", bold=True)
        elif status.upper() == "FAIL":
            status_cell.fill = PatternFill(start_color="FFC7CE", end_color="FFC7CE", fill_type="solid")
            status_cell.font = Font(color="9C0006", bold=True)
        else:  # INFO
            status_cell.fill = PatternFill(start_color="FFEB9C", end_color="FFEB9C", fill_type="solid")
            status_cell.font = Font(color="9C6500", bold=True)

        # Align all cells in the row
        for col in range(1, 5):
            self.sheet.cell(row=last_row, column=col).alignment = Alignment(
                horizontal="left", vertical="center", wrap_text=True
            )

        self.wb.save(self.filename)

    def safe_action(self, driver, action_desc, func):
        """Runs a Selenium action and handles alerts and other exceptions cleanly."""
        try:
            func()
            self.log("PASS", action_desc, "Action completed successfully.")
        except UnexpectedAlertPresentException:
            try:
                alert = driver.switch_to.alert
                alert_text = alert.text
                self.log("FAIL", action_desc, f"Unexpected Alert: {alert_text}")
                alert.accept()
            except NoAlertPresentException:
                self.log("FAIL", action_desc, "Alert appeared but no text found.")
        except Exception as e:
            # Log only the error message, not the full stack trace.
            self.log("FAIL", action_desc, f"Error: {type(e).__name__} - {str(e)}")

# --- Main Script Execution ---
# Initialize the logger
logger = ExcelLogger()
logger.log("INFO", "Script started.", "Initializing Excel log and WebDriver.")

# launch chrome
try:
    driver = webdriver.Chrome()
    driver.maximize_window()
    wait = WebDriverWait(driver, 10)
    logger.log("INFO", "Chrome browser launched.", "Maximized window and set WebDriverWait.")
except Exception as e:
    logger.log("FAIL", "Launch browser.", f"Error: {str(e)}")
    exit()

# go to url
logger.safe_action(driver, "Navigate to URL", lambda: driver.get("http://code-server-app/CSS/LoginPage.aspx?"))
time.sleep(2)

# login button
logger.safe_action(driver, "Enter username", lambda: wait.until(EC.presence_of_element_located((By.NAME, "txtUsername"))).send_keys("admin"))
logger.safe_action(driver, "Enter password", lambda: wait.until(EC.presence_of_element_located((By.ID, "txtPassword"))).send_keys('admin'))
logger.safe_action(driver, "Click Login", lambda: wait.until(EC.element_to_be_clickable((By.ID, "btnLogin"))).click())
time.sleep(2)

# Region Management
logger.safe_action(driver, "Navigate to Region", lambda: wait.until(EC.presence_of_element_located((By.XPATH, "//a[@href='/CSS/Region.aspx']"))).click())
logger.safe_action(driver, "Click Add Region", lambda: wait.until(EC.element_to_be_clickable((By.ID, "btnAddRegion"))).click())
time.sleep(1)

logger.safe_action(driver, "Enter region name", lambda: driver.find_element(By.ID, "MRegion").send_keys("Jammu"))
logger.safe_action(driver, "Enter description", lambda: driver.find_element(By.ID, "MDescription").send_keys("Test Description"))
time.sleep(2)

checkbox = driver.find_element(By.ID, "MisActive")
if not checkbox.is_selected():
    logger.safe_action(driver, "Check 'Is Active' checkbox", lambda: checkbox.click())
else:
    logger.log("INFO", "'Is Active' checkbox is already checked.")

logger.safe_action(driver, "Save New Region", lambda: driver.find_element(By.ID, "ContentPlaceHolder1_btnSave").click())
time.sleep(2)
logger.log("INFO", "Region 'Jammu' added successfully.", "Log entry for newly added region.")

# Edit Region
logger.safe_action(driver, "Click Edit", lambda: wait.until(EC.element_to_be_clickable((By.XPATH, "(//i[contains(@class,'mdi-square-edit-outline')])[1]"))).click())
logger.safe_action(driver, "Clear and enter new region name", lambda: wait.until(EC.visibility_of_element_located((By.ID, "MRegion"))).clear())
logger.safe_action(driver, "Clear and enter new region name", lambda: driver.find_element(By.ID, "MRegion").send_keys("Jammu & Kashmir"))

logger.safe_action(driver, "Save Edited Region", lambda: driver.find_element(By.ID, "ContentPlaceHolder1_btnSave").click())
time.sleep(2)

# Branch Management
logger.safe_action(driver, "Navigate to Branch", lambda: wait.until(EC.presence_of_element_located((By.XPATH, "//a[@href='/CSS/Branch.aspx']"))).click())
time.sleep(1)
logger.safe_action(driver, "Click Add Branch", lambda: wait.until(EC.presence_of_element_located((By.ID, "btnAddBranch"))).click())
time.sleep(1)

logger.safe_action(driver, "Select Region 'Vidarbha'", lambda: Select(driver.find_element(By.ID, "ddlRegion")).select_by_visible_text("Vidarbha"))
logger.safe_action(driver, "Enter Branch Name", lambda: wait.until(EC.presence_of_element_located((By.ID, "txtBranchName"))).send_keys("Inox"))
logger.safe_action(driver, "Enter City", lambda: wait.until(EC.presence_of_element_located((By.ID, "txtCityName"))).send_keys("Mumbai"))
logger.safe_action(driver, "Select size radio button", lambda: wait.until(EC.presence_of_element_located((By.ID, "RBBySize"))).click())
logger.safe_action(driver, "Enter size", lambda: wait.until(EC.presence_of_element_located((By.ID, "txtBySize"))).send_keys("2"))
logger.safe_action(driver, "Check 'Is Active' for branch", lambda: wait.until(EC.presence_of_element_located((By.ID, "ChkIsActive"))).click())
logger.safe_action(driver, "Save Branch", lambda: wait.until(EC.presence_of_element_located((By.ID, "ContentPlaceHolder1_btnSave"))).click())
time.sleep(1)

# FTP Configuration
logger.safe_action(driver, "Navigate to FTP", lambda: wait.until(EC.presence_of_element_located((By.XPATH, "(//i[contains(@class, 'mdi-server')])[3]"))).click())
time.sleep(1)
logger.safe_action(driver, "Enter FTP username", lambda: wait.until(EC.presence_of_element_located((By.ID, "txtusername"))).clear())
logger.safe_action(driver, "Enter FTP username", lambda: driver.find_element(By.ID, "txtusername").send_keys("dae123"))
logger.safe_action(driver, "Enter server name", lambda: wait.until(EC.presence_of_element_located((By.ID, "txtservername"))).clear())
logger.safe_action(driver, "Enter server name", lambda: driver.find_element(By.ID, "txtservername").send_keys("code-server-app"))
logger.safe_action(driver, "Enter server password", lambda: wait.until(EC.presence_of_element_located((By.ID, "txtpassword"))).clear())
logger.safe_action(driver, "Enter server password", lambda: driver.find_element(By.ID, "txtpassword").send_keys("P@ssw0rd786"))
logger.safe_action(driver, "Enter database name", lambda: wait.until(EC.presence_of_element_located((By.NAME, "ctl00$ContentPlaceHolder1$txtdatabasename"))).clear())
logger.safe_action(driver, "Enter database name", lambda: driver.find_element(By.NAME, "ctl00$ContentPlaceHolder1$txtdatabasename").send_keys("code-server-app"))
time.sleep(0.5)

logger.safe_action(driver, "Check central FTP", lambda: wait.until(EC.presence_of_element_located((By.ID, "ContentPlaceHolder1_centeralcheck"))).click())
logger.safe_action(driver, "Enter central FTP server name", lambda: wait.until(EC.presence_of_element_located((By.NAME, "ctl00$ContentPlaceHolder1$txtcenterftpname"))).clear())
logger.safe_action(driver, "Enter central FTP server name", lambda: driver.find_element(By.NAME, "ctl00$ContentPlaceHolder1$txtcenterftpname").send_keys("DAE"))
logger.safe_action(driver, "Enter central FTP username", lambda: wait.until(EC.presence_of_element_located((By.ID, "txtcenteruser"))).clear())
logger.safe_action(driver, "Enter central FTP username", lambda: driver.find_element(By.ID, "txtcenteruser").send_keys("dae123"))
logger.safe_action(driver, "Enter central FTP password", lambda: wait.until(EC.presence_of_element_located((By.ID, "txtcenterpassword"))).clear())
logger.safe_action(driver, "Enter central FTP password", lambda: driver.find_element(By.ID, "txtcenterpassword").send_keys("P@ssw0rd786"))
logger.safe_action(driver, "Add FTP", lambda: wait.until(EC.presence_of_element_located((By.ID, "ContentPlaceHolder1_Add"))).click())
time.sleep(2)

# Media
logger.safe_action(driver, "Navigate to Media", lambda: wait.until(EC.presence_of_element_located((By.XPATH, "//a[@href='/CSS/Media.aspx']"))).click())
time.sleep(0.5)
logger.safe_action(driver, "Click 'Add Media'", lambda: wait.until(EC.presence_of_element_located((By.ID, "exampleModal"))).click())
file_path = "E:\\Jude\\reebok.jpeg"
logger.safe_action(driver, "Upload file", lambda: wait.until(EC.presence_of_element_located((By.ID, "FileUpload1"))).send_keys(file_path))
time.sleep(0.5)
logger.safe_action(driver, "Enter caption", lambda: wait.until(EC.presence_of_element_located((By.ID, "Caption"))).send_keys("Wallpaper"))
logger.safe_action(driver, "Enter duration", lambda: wait.until(EC.presence_of_element_located((By.ID, "txtDuration"))).send_keys("10"))
logger.safe_action(driver, "Enter start date", lambda: wait.until(EC.presence_of_element_located((By.ID, "txtStartDate"))).send_keys("14-08-2025"))
logger.safe_action(driver, "Enter end date", lambda: wait.until(EC.presence_of_element_located((By.ID, "TxtEndDate"))).send_keys("16-08-2025"))
is_act = wait.until(EC.presence_of_element_located((By.ID, "ChkIsActive")))
if not is_act.is_selected():
    logger.safe_action(driver, "Check 'Is Active' for media", lambda: is_act.click())
logger.safe_action(driver, "Save media", lambda: wait.until(EC.presence_of_element_located((By.ID, "ContentPlaceHolder1_Button1"))).click())
time.sleep(5)

# Playlist
logger.safe_action(driver, "Navigate to Playlist", lambda: wait.until(EC.presence_of_element_located((By.XPATH, "//a[@href='/CSS/Playlist.aspx']"))).click())
time.sleep(2)
logger.safe_action(driver, "Click eye icon", lambda: wait.until(EC.element_to_be_clickable((By.XPATH, "(//i[contains(@class, 'mdi-eye')])[1]"))).click())
time.sleep(2)

action = ActionChains(driver)
drag = wait.until(EC.presence_of_element_located((By.ID, "media_135")))
drop = wait.until(EC.presence_of_element_located((By.ID, "dropzone")))
drag1 = wait.until(EC.presence_of_element_located((By.ID, "media_136")))
drop1 = wait.until(EC.presence_of_element_located((By.ID, "dropzone")))

logger.safe_action(driver, "Drag and drop media 135", lambda: action.drag_and_drop(drag, drop).perform())
logger.safe_action(driver, "Drag and drop media 136", lambda: action.drag_and_drop(drag1, drop1).perform())
time.sleep(2)

driver.execute_script("window.scrollBy(0, 200);")
logger.safe_action(driver, "Scroll down", lambda: None)

logger.safe_action(driver, "Add slot", lambda: wait.until(EC.presence_of_element_located((By.ID, "ContentPlaceHolder1_btnOpenSlotAdd"))).click())
slot_name = wait.until(EC.visibility_of_element_located((By.ID, "slotname")))
logger.safe_action(driver, "Scroll to slot name", lambda: driver.execute_script("arguments[0].scrollIntoView({block:'center'});", slot_name))
logger.safe_action(driver, "Enter slot name", lambda: slot_name.send_keys("Temporary"))
logger.safe_action(driver, "Enter start time", lambda: wait.until(EC.presence_of_element_located((By.ID, "txtStartDateslot"))).send_keys("12.45"))
logger.safe_action(driver, "Enter end time", lambda: wait.until(EC.presence_of_element_located((By.ID, "txtEndDateslot"))).send_keys("14.45"))
slot_active = wait.until(EC.presence_of_element_located((By.ID, "CheckBox")))
if not slot_active.is_selected():
    logger.safe_action(driver, "Check 'Is Active' for slot", lambda: slot_active.click())
logger.safe_action(driver, "Add slot", lambda: wait.until(EC.presence_of_element_located((By.ID, "addslot"))).click())
time.sleep(2)
logger.safe_action(driver, "Close slot dialog", lambda: wait.until(EC.presence_of_element_located((By.ID, "ContentPlaceHolder1_clear"))).click())
time.sleep(5)

# Screen
logger.safe_action(driver, "Navigate to Screen", lambda: wait.until(EC.presence_of_element_located((By.XPATH, "//i[@class='mdi mdi-monitor']"))).click())
time.sleep(2)
logger.safe_action(driver, "Click 'Add Screen'", lambda: wait.until(EC.presence_of_element_located((By.ID, "btnAddScreen"))).click())
time.sleep(0.8)
screen_name = wait.until(EC.presence_of_element_located((By.ID, "txtScreen")))
logger.safe_action(driver, "Enter screen name", lambda: screen_name.clear())
logger.safe_action(driver, "Enter screen name", lambda: screen_name.send_keys("Screen 151"))
time.sleep(1.5)
logger.safe_action(driver, "Select playlist", lambda: Select(wait.until(EC.element_to_be_clickable((By.ID, "ddlPlaylist")))).select_by_visible_text("Boisar"))
logger.safe_action(driver, "Select screen branch", lambda: Select(wait.until(EC.element_to_be_clickable((By.ID, "ddlBranchlist")))).select_by_visible_text("Growels 101"))
logger.safe_action(driver, "Save screen", lambda: wait.until(EC.presence_of_element_located((By.ID, "ContentPlaceHolder1_Button1"))).click())
time.sleep(2)

# Ticker
logger.safe_action(driver, "Navigate to Ticker", lambda: wait.until(EC.presence_of_element_located((By.XPATH, "//a[@href='/CSS/TickerMaster.aspx']"))).click())
time.sleep(0.5)
logger.safe_action(driver, "Click 'Add Ticker'", lambda: wait.until(EC.presence_of_element_located((By.ID, "btnAddPlaylist"))).click())
time.sleep(0.5)
logger.safe_action(driver, "Enter ticker tag", lambda: wait.until(EC.presence_of_element_located((By.ID, "txttickertag"))).send_keys("Ticker Infiniti"))
logger.safe_action(driver, "Enter ticker text", lambda: wait.until(EC.presence_of_element_located((By.ID, "txttickertext"))).send_keys("WELCOME TO INFINITI  "))
logger.safe_action(driver, "Select font name", lambda: Select(wait.until(EC.presence_of_element_located((By.ID, "ddlFontName")))).select_by_visible_text("Calibri"))
logger.safe_action(driver, "Enter font size", lambda: wait.until(EC.presence_of_element_located((By.ID, "txtFontName"))).send_keys("2"))
logger.safe_action(driver, "Enter font speed", lambda: wait.until(EC.presence_of_element_located((By.ID, "txtSpeed"))).send_keys("2"))
direction = wait.until(EC.presence_of_element_located((By.ID, "RdoRTL")))
if not direction.is_selected():
    logger.safe_action(driver, "Click direction", lambda: direction.click())
ticker_active = wait.until(EC.presence_of_element_located((By.ID, "ChkIsActive")))
if not ticker_active.is_selected():
    logger.safe_action(driver, "Check 'Is Active' for ticker", lambda: ticker_active.click())
logger.safe_action(driver, "Save Ticker", lambda: wait.until(EC.presence_of_element_located((By.ID, "ContentPlaceHolder1_btnSave"))).click())
time.sleep(1)

# Profile and Logout
logger.safe_action(driver, "Click profile", lambda: driver.find_element(By.XPATH, "//a[@role='button' and @data-bs-toggle='dropdown' and contains(@class, 'nav-user')]").click())
time.sleep(1)
logger.safe_action(driver, "Click Logout", lambda: driver.find_element(By.XPATH, "//a[@href='?action=logout']").click())
time.sleep(2)
logger.safe_action(driver, "Accept alert", lambda: driver.switch_to.alert.accept())

# Final cleanup
logger.log("INFO", "Script finished.", "Closing the browser.")
driver.quit()