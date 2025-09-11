import time
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

from selenium import webdriver
from selenium.webdriver.common.by import By
from selenium.webdriver.support.ui import Select
from selenium.webdriver.common.action_chains import ActionChains
from selenium.webdriver.support.ui import WebDriverWait
from selenium.webdriver.support import expected_conditions as EC


# --- Excel Logging Setup ---
def setup_excel_log(filename):
    """
    Creates a new Excel workbook with a log sheet and formatted headers.
    """
    try:
        workbook = openpyxl.Workbook()
        sheet = workbook.active
        sheet.title = "Chainvest Automation Log"

        headers = ["Timestamp", "Step Description", "Status", "Details"]
        sheet.append(headers)

        header_font = Font(bold=True, color="FFFFFF")
        header_fill = PatternFill(start_color="4F81BD", end_color="4F81BD", fill_type="solid")
        for col_num, header in enumerate(headers, 1):
            cell = sheet.cell(row=1, column=col_num)
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = Alignment(horizontal='center')

        sheet.column_dimensions['B'].width = 60
        sheet.column_dimensions['C'].width = 15
        sheet.column_dimensions['D'].width = 80

        workbook.save(filename)
        return workbook, sheet
    except Exception as e:
        print(f"Error setting up Excel file: {e}")
        return None, None


def log_step(sheet, description, status, details=""):
    """
    Logs a step's result to the Excel sheet with color-coding.
    """
    try:
        timestamp = time.strftime("%Y-%m-%d %H:%M:%S")
        log_data = [timestamp, description, status, details]
        sheet.append(log_data)

        row = sheet.max_row
        status_cell = sheet.cell(row=row, column=3)
        description_cell = sheet.cell(row=row, column=2)
        details_cell = sheet.cell(row=row, column=4)

        if status.upper() == "PASS":
            status_cell.fill = PatternFill(start_color="C6EFCE", end_color="C6EFCE", fill_type="solid")
            status_cell.font = Font(color="006100")
        elif status.upper() == "FAIL":
            status_cell.fill = PatternFill(start_color="FFC7CE", end_color="FFC7CE", fill_type="solid")
            status_cell.font = Font(color="9C0006")
        elif status.upper() == "INFO":
            status_cell.fill = PatternFill(start_color="B9CDE5", end_color="B9CDE5", fill_type="solid")
            status_cell.font = Font(color="0000FF")

        description_cell.alignment = Alignment(wrap_text=True, vertical='top')
        details_cell.alignment = Alignment(wrap_text=True, vertical='top')

    except Exception as e:
        print(f"Error logging step: {e}")


# --- Main Script Execution ---
lof_folder =  r"C:\Users\manoj\Desktop\Logs\ChainvestLogs"
log_filename = "Chainvest_Automation_Log_2608.xlsx"
workbook, sheet = setup_excel_log(log_filename)
if not workbook:
    exit()

try:
    # --- Setting up browser ---
    log_step(sheet, "Starting browser setup", "INFO")
    options = webdriver.ChromeOptions()
    options.add_argument("--incognito")
    driver = webdriver.Chrome(options=options)
    driver.maximize_window()
    wait = WebDriverWait(driver, 20)  # Increased wait time for stability
    log_step(sheet, "Browser setup completed successfully", "PASS")

    # --- Navigating to login page and logging in ---
    log_step(sheet, "Attempting login", "INFO", "Navigating to login page and entering credentials.")
    driver.get("http://code-server-app/Chainvest/BO/login")
    wait.until(EC.presence_of_element_located((By.ID, "Email"))).send_keys("admin@chainvest.com")
    driver.find_element(By.ID, "Password").send_keys("P@ssw0rd")
    driver.find_element(By.XPATH, "//button[normalize-space(text())='Sign in']").click()
    wait.until(EC.url_to_be("http://code-server-app/Chainvest/BO/dashboard"))
    log_step(sheet, "Login successful", "PASS", "Successfully logged in as admin.")

    # --- Adding a new business (DMart) ---
    log_step(sheet, "Adding a new business: DMart", "INFO")
    try:
        wait.until(EC.element_to_be_clickable((By.XPATH, "//span[normalize-space(text())='Business']"))).click()
        wait.until(EC.element_to_be_clickable((By.XPATH, "//button[contains(., 'Add Business')]"))).click()
        driver.find_element(By.ID, "Name").send_keys("DMart")
        select = Select(driver.find_element(By.XPATH, "//select[@id='BusinessTypeID']"))
        select.select_by_visible_text("Retailer")
        driver.find_element(By.XPATH, "//input[@role='combobox']").send_keys("Retail")
        wait.until(EC.element_to_be_clickable((By.XPATH, "//span[normalize-space()='Retail']"))).click()
        driver.find_element(By.XPATH, "//input[@id='Logo']").send_keys(r"C:\Users\manoj\Desktop\SampleImages\dmart.png")
        driver.find_element(By.ID, "ContactPersonName").send_keys("DMart Admin")
        driver.find_element(By.ID, "EmailID").send_keys("dmart@sample.com")
        driver.find_element(By.ID, "MobileNo1").send_keys("99911212111")
        driver.find_element(By.ID, "MobileNo2").send_keys("9786411087")
        driver.find_element(By.ID, "WebsiteURL").send_keys("https://www.dmartindia.com/")
        driver.find_element(By.ID, "BusinessLicenseNo").send_keys("A12345")
        driver.find_element(By.ID, "VAT").send_keys("IE1234567T")
        driver.find_element(By.ID, "Description").send_keys("Test description")
        driver.find_element(By.XPATH, "//span[normalize-space()='Save & Next']").click()

        doc1_input = wait.until(
            EC.presence_of_element_located((By.XPATH, "//button[normalize-space(text())='Upload Document']")))
        doc1_input.send_keys(r"C:\Users\manoj\Desktop\dmart_sample.jpg")
        driver.find_element(By.XPATH, "//button[span[normalize-space()='Save & Next']]").click()

        address1 = wait.until(
            EC.presence_of_element_located((By.XPATH, "//input[@type='text' and @formcontrolname='Address1']")))
        address1.send_keys("Bangur Nagar")
        driver.find_element(By.XPATH, "//input[@type='text' and @formcontrolname='Address2']").send_keys(
            "Goregaon West")
        select_country = Select(driver.find_element(By.XPATH, "//select[@formcontrolname='CountryID']"))
        select_country.select_by_visible_text("Ethiopia")
        wait.until(EC.visibility_of_element_located(
            (By.XPATH, "//select[@formcontrolname='StateID']/option[text()='Addis Ababa']")))
        select_state = Select(
            wait.until(EC.element_to_be_clickable((By.XPATH, "//select[@formcontrolname='StateID']"))))
        select_state.select_by_visible_text("Addis Ababa")
        wait.until(EC.visibility_of_element_located(
            (By.XPATH, "//select[@formcontrolname='CityID']/option[normalize-space()='Addis Ababa']")))
        select_city = Select(wait.until(EC.element_to_be_clickable((By.XPATH, "//select[@formcontrolname='CityID']"))))
        select_city.select_by_visible_text("Addis Ababa")
        driver.find_element(By.XPATH, "//input[@formcontrolname='Zip']").send_keys("400123")
        driver.find_element(By.XPATH, "//input[@formcontrolname='Landmark']").send_keys("Addis")
        driver.find_element(By.XPATH, "//input[@formcontrolname='Latitude']").send_keys("40.12")
        driver.find_element(By.XPATH, "//input[@formcontrolname='Longitude']").send_keys("35.98")
        driver.find_element(By.XPATH, "//span[normalize-space()='Save & Next']").click()

        wait.until(EC.presence_of_element_located((By.XPATH, "//input[@formcontrolname='BankName']"))).send_keys("HDFC")
        driver.find_element(By.XPATH, "//input[@formcontrolname='AccountHolderName']").send_keys("DMart")
        driver.find_element(By.XPATH, "//input[@formcontrolname= 'AccountNumber']").send_keys("2874271980167")
        driver.find_element(By.XPATH, "//input[@formcontrolname= 'IFSC']").send_keys("HDFC00000504")
        driver.find_element(By.XPATH, "//input[@formcontrolname= 'BankBranchName']").send_keys("HDFC Addis")
        driver.find_element(By.XPATH, "//span[normalize-space()='Save & Complete']").click()
        log_step(sheet, "New business 'DMart' added successfully", "PASS")
    except Exception as e:
        log_step(sheet, "Failed to add new business 'DMart'", "FAIL", f"Error: {e}")

    # --- Master Module (Business Type, Roles, Payment Type, Payment Status, Order Status) ---
    log_step(sheet, "Starting Master module configurations", "INFO")
    try:
        # Business Type
        wait.until(EC.presence_of_element_located((By.XPATH, "//span[normalize-space() = 'Master']"))).click()
        wait.until(EC.presence_of_element_located((By.XPATH, "//a[@href='/Chainvest/BO/business-type']"))).click()
        wait.until(
            EC.presence_of_element_located((By.XPATH, "//button[normalize-space()= 'Add Business Type']"))).click()
        driver.find_element(By.XPATH, "//input[@formcontrolname='BusinessTypeName']").send_keys("Agent")
        driver.find_element(By.XPATH, "//textarea[@formcontrolname='Description']").send_keys(
            "Acts on behalf of a company to sell or promote products/services, usually earning a commission.")
        driver.find_element(By.XPATH, "//button[normalize-space(text())='Add']").click()
        log_step(sheet, "New Business Type 'Agent' added", "PASS")

        # Roles
        driver.find_element(By.XPATH, "//a[@href='/Chainvest/BO/roles']").click()
        driver.find_element(By.XPATH, "//button[normalize-space(text())='Add Role']").click()
        select = Select(driver.find_element(By.XPATH, "//select[@formcontrolname='PrimaryRoleID']"))
        select.select_by_visible_text("EndUser")
        driver.find_element(By.XPATH, "//input[@formcontrolname='Name']").send_keys("End User")
        driver.find_element(By.XPATH, "//textarea[@formcontrolname='Description']").send_keys("Application User")
        driver.find_element(By.XPATH, "//button[normalize-space(text())='Add']").click()
        log_step(sheet, "New Role 'End User' added", "PASS")

        # Payment Type
        driver.find_element(By.XPATH, "//a[@href='/Chainvest/BO/paymentType']").click()
        driver.find_element(By.XPATH, "//button[normalize-space(text())='Add Payment Type']").click()
        driver.find_element(By.XPATH, "//input[@formcontrolname='PaymentTypeName']").send_keys("QR Code")
        driver.find_element(By.XPATH, "//textarea[@formcontrolname='PaymentTypeDescription']").send_keys(
            "Scan an QR code for faster payment")
        driver.find_element(By.XPATH, "//button[normalize-space(text())='Add']").click()
        log_step(sheet, "New Payment Type 'QR Code' added", "PASS")

        # Payment Status
        driver.find_element(By.XPATH, "//a[@href='/Chainvest/BO/paymentStatus']").click()
        driver.find_element(By.XPATH, "//button[normalize-space(text())='Add Payment Status']").click()
        driver.find_element(By.XPATH, "//input[@formcontrolname='StatusName']").send_keys("Initiated")
        driver.find_element(By.XPATH, "//button[normalize-space(text())='Add']").click()
        log_step(sheet, "New Payment Status 'Initiated' added", "PASS")

        # Order Status
        driver.find_element(By.XPATH, "//a[@href='/Chainvest/BO/order-status']").click()
        driver.find_element(By.XPATH, "//button[normalize-space(text())='Add Order Status']").click()
        driver.find_element(By.ID, "OrderStatusName").send_keys("Pending")
        driver.find_element(By.ID, "OrderStatusDescription").send_keys("Order has been placed but not yet processed.")
        driver.find_element(By.XPATH, "//button[normalize-space(text())='Add']").click()
        log_step(sheet, "New Order Status 'Pending' added", "PASS")
    except Exception as e:
        log_step(sheet, "Failed to configure Master modules", "FAIL", f"Error: {e}")

    # --- Loan Setting Module ---
    log_step(sheet, "Starting Loan Setting module configurations", "INFO")
    try:
        # Loan Purpose
        wait.until(EC.presence_of_element_located((By.XPATH, "//span[normalize-space(text())='Loan Setting']"))).click()
        driver.find_element(By.XPATH, "//a[@href='/Chainvest/BO/loan-purpose']").click()
        driver.find_element(By.XPATH, "//button[normalize-space(text())='Add Loan Purpose']").click()
        driver.find_element(By.ID, "LoanPurposeName").send_keys("Logistics")
        driver.find_element(By.ID, "LoanPurposeDescription").send_keys("Fleet expansion or delivery service payments")
        driver.find_element(By.XPATH, "//input[@formcontrolname='InterestFrom']").send_keys("8")
        driver.find_element(By.XPATH, "//input[@formcontrolname='InterestTO']").send_keys("20")
        driver.find_element(By.ID, "name").send_keys("200000")
        driver.find_element(By.ID, "LoanAmountTO").send_keys("800000")
        driver.find_element(By.XPATH, "//button[normalize-space(text())='Add']").click()
        log_step(sheet, "New Loan Purpose 'Logistics' added", "PASS")

        # Loan Status
        driver.find_element(By.XPATH, "//a[@href='/Chainvest/BO/loan-status']").click()
        driver.find_element(By.XPATH, "//button[normalize-space(text())='Add Loan Status']").click()
        driver.find_element(By.ID, "LoanStatusName").send_keys("Overdue")
        driver.find_element(By.ID, "LoanStatusDescription").send_keys("Payment is delayed beyond the due date.")
        driver.find_element(By.XPATH, "//button[normalize-space(text())='Add']").click()
        log_step(sheet, "New Loan Status 'Overdue' added", "PASS")
    except Exception as e:
        log_step(sheet, "Failed to configure Loan Setting modules", "FAIL", f"Error: {e}")

    # --- Credit Score Module ---
    log_step(sheet, "Starting Credit Score module configurations", "INFO")
    try:
        # Add Rule
        wait.until(EC.presence_of_element_located((By.XPATH, "//span[normalize-space(text())='Credit Score']"))).click()
        driver.find_element(By.XPATH, "//a[@href='/Chainvest/BO/credit-score']").click()
        scrollable_module = driver.find_element(By.CLASS_NAME, "overflow-y-scroll")
        driver.execute_script("arguments[0].scrollTop = arguments[0].scrollHeight;", scrollable_module)
        time.sleep(1)
        driver.execute_script("arguments[0].scrollTop = 0;", scrollable_module)
        wait.until(EC.element_to_be_clickable((By.XPATH, "//button[text()=' Add Rule ']"))).click()
        wait.until(EC.presence_of_element_located((By.ID, "RuleName"))).send_keys("Credit Utilization")
        driver.find_element(By.ID, "Description").send_keys("Credit utilization ratio between 10% and 30%")
        driver.find_element(By.ID, "ScoreImpact").send_keys("10")
        prio = driver.find_element(By.ID, "Priority")
        prio.clear()
        prio.send_keys("4")
        driver.find_element(By.XPATH, "//button[normalize-space(text())='Create Rule']").click()
        log_step(sheet, "New Credit Rule 'Credit Utilization' added", "PASS")

        # View Score
        wait.until(EC.presence_of_element_located((By.XPATH, "//a[@href='/Chainvest/BO/credit-scorescreen']"))).click()
        wait.until(EC.presence_of_element_located((By.XPATH, "//li[text()=' InnovatePro Fabrications ']"))).click()
        log_step(sheet, "Credit score for 'InnovatePro Fabrications' viewed", "PASS")
    except Exception as e:
        log_step(sheet, "Failed to configure Credit Score modules", "FAIL", f"Error: {e}")

    # --- Product Module ---
    log_step(sheet, "Starting Product module configurations", "INFO")
    try:
        # Bulk Upload
        wait.until(EC.presence_of_element_located((By.XPATH, "//span[normalize-space()='Product']"))).click()
        driver.find_element(By.XPATH, "//a[@href='/Chainvest/BO/bulk-upload']").click()
        wait.until(EC.presence_of_element_located((By.ID, "fileInput"))).send_keys(
            r"C:\Users\manoj\Downloads\chainvestupload.xlsx")
        driver.find_element(By.XPATH, "//button[normalize-space()='Upload File']").click()
        log_step(sheet, "Product Bulk Upload completed", "PASS")

        # Product Category
        driver.find_element(By.XPATH, "//a[@href='/Chainvest/BO/product-category']").click()
        wait.until(EC.presence_of_element_located(
            (By.XPATH, "//button[normalize-space(text())='Add Product Category']"))).click()
        wait.until(EC.presence_of_element_located((By.XPATH, "//input[@type='file']"))).send_keys(
            r"C:\Users\manoj\Downloads\earbuds.jpeg")
        driver.find_element(By.XPATH, "(//input[@type='file'])[2]").send_keys(r"C:\Users\manoj\Downloads\earbuds.jpeg")

        # Enter details for each language
        languages = {
            "Hindi": ("Earbuds", "Earbuds"),
            "Afan Oromo": ("Earbuds", "Earbuds"),
            "Amharic": ("Earbuds", "Earbuds"),
            "English": ("Earbuds", "Earbuds")
        }
        for lang, (name, desc) in languages.items():
            driver.find_element(By.XPATH,
                                f"//tr[td[text()='{lang}']]//input[@formcontrolname='ProductCategoryShortName']").send_keys(
                name)
            driver.find_element(By.XPATH,
                                f"//tr[td[text()='{lang}']]//input[@placeholder='Enter Description']").send_keys(desc)

        driver.find_element(By.XPATH, "//button[normalize-space()='Add']").click()
        log_step(sheet, "New Product Category 'Earbuds' added", "PASS")

        # Sub-Category
        wait.until(EC.element_to_be_clickable((By.XPATH, "//a[@href='/Chainvest/BO/product-subcategory']"))).click()
        wait.until(EC.presence_of_element_located((By.XPATH, "//button[text()='Add Product SubCategory']"))).click()
        select_sub_category = Select(
            wait.until(EC.presence_of_element_located((By.XPATH, "//select[@formcontrolname='ProductCategoryID']"))))
        select_sub_category.select_by_visible_text("Earbuds")
        driver.find_element(By.XPATH, "(//input[@type='file'])[1]").send_keys(r"C:\Users\manoj\Downloads\earbuds.jpeg")
        driver.find_element(By.XPATH, "(//input[@type='file'])[2]").send_keys(r"C:\Users\manoj\Downloads\earbuds.jpeg")

        languages_sub = {
            "Hindi": ("Wireless earbuds", "Wireless earbuds"),
            "Afan Oromo": ("Wireless earbuds", "Wireless earbuds"),
            "Amharic": ("Wireless earbuds", "Wireless earbuds"),
            "English": ("Wireless Earbuds", "Wireless earbuds")
        }
        for lang, (name, desc) in languages_sub.items():
            driver.find_element(By.XPATH,
                                f"//tr[td[text()='{lang}']]//input[@formcontrolname='ProductSubCategoryShortName']").send_keys(
                name)
            driver.find_element(By.XPATH,
                                f"//tr[td[text()='{lang}']]//input[@placeholder='Enter Description']").send_keys(desc)

        driver.find_element(By.XPATH, "//button[normalize-space()='Add']").click()
        log_step(sheet, "New Sub-Category 'Wireless Earbuds' added", "PASS")
    except Exception as e:
        log_step(sheet, "Failed to configure Product modules", "FAIL", f"Error: {e}")

    # --- Orders Module ---
    log_step(sheet, "Viewing Orders", "INFO")
    try:
        wait.until(EC.presence_of_element_located((By.XPATH, "//span[normalize-space(text())='Orders']"))).click()
        wait.until(EC.presence_of_element_located((By.XPATH, "//button[normalize-space(text())='View']"))).click()
        log_step(sheet, "Order view successful", "PASS")
    except Exception as e:
        log_step(sheet, "Failed to view Orders", "FAIL", f"Error: {e}")

    # --- Users Module ---
    log_step(sheet, "Adding a new user", "INFO")
    try:
        wait.until(EC.presence_of_element_located((By.XPATH, "//span[normalize-space(text())='Users']"))).click()
        wait.until(EC.presence_of_element_located((By.XPATH, "//button[normalize-space(text())='Add User']"))).click()
        select_company = Select(
            wait.until(EC.presence_of_element_located((By.XPATH, "//select[@formcontrolname = 'CompanyID']"))))
        select_company.select_by_visible_text("Digital Automation Enterprises")
        driver.find_element(By.XPATH, "//input[@formcontrolname = 'Email']").send_keys("admin@dae.com")
        driver.find_element(By.XPATH, "//input[@formcontrolname = 'PasswordHash']").send_keys("P@ssw0rd")
        country_select = Select(driver.find_element(By.XPATH, "//select[@formcontrolname = 'CountryID']"))
        country_select.select_by_visible_text("Ethiopia")
        driver.find_element(By.XPATH, "//input[@formcontrolname = 'Phone']").send_keys("9874615860")
        role_select = Select(driver.find_element(By.XPATH, "//select[@formcontrolname = 'RoleID']"))
        role_select.select_by_visible_text("Admin")
        driver.find_element(By.XPATH, "//button[normalize-space(text())='Add']").click()
        log_step(sheet, "New user 'admin@dae.com' added", "PASS")
    except Exception as e:
        log_step(sheet, "Failed to add new user", "FAIL", f"Error: {e}")

    # --- Reports Module ---
    log_step(sheet, "Generating Reports", "INFO")
    try:
        wait.until(EC.presence_of_element_located((By.XPATH, "//span[normalize-space(text())='Reports']"))).click()

        # Order Report
        wait.until(EC.presence_of_element_located((By.XPATH, "//a[@href='/Chainvest/BO/fulfillment-report']"))).click()
        wait.until(EC.presence_of_element_located((By.XPATH, "//input[@formcontrolname='orderDateFrom']"))).send_keys(
            "01-08-2025")
        driver.find_element(By.XPATH, "//input[@formcontrolname='orderDateTo']").send_keys("20-08-2025")
        driver.find_element(By.XPATH, "//button[normalize-space(text())='Search']").click()
        time.sleep(2)
        driver.find_element(By.XPATH, "//button[normalize-space(text())='Clear']").click()
        log_step(sheet, "Order Report generated successfully", "PASS")

        # Bank and Supply Chain Report
        driver.find_element(By.XPATH, "//a[@href='/Chainvest/BO/bankandsupplychain-report']").click()
        log_step(sheet, "Bank and Supply Chain Report viewed", "PASS")

        # Financial Report
        driver.find_element(By.XPATH, "//a[@href='/Chainvest/BO/financial-report']").click()
        log_step(sheet, "Financial Report viewed", "PASS")
    except Exception as e:
        log_step(sheet, "Failed to generate Reports", "FAIL", f"Error: {e}")

    # --- Loan Module View ---
    log_step(sheet, "Viewing Loan Details", "INFO")
    try:
        loan = driver.find_element(By.XPATH, "//span[normalize-space(text())='Loan']")
        loan.click()
        wait.until(EC.element_to_be_clickable((By.XPATH, "//button[normalize-space()='View']"))).click()
        driver.execute_script("window.scrollTo(0, document.body.scrollHeight);")
        time.sleep(1)
        driver.execute_script("window.scrollTo(0, 0);")
        log_step(sheet, "Loan details viewed successfully", "PASS")
    except Exception as e:
        log_step(sheet, "Failed to view Loan details", "FAIL", f"Error: {e}")

    log_step(sheet, "Automation script completed successfully.", "INFO")

except Exception as e:
    log_step(sheet, "An unexpected error occurred", "FAIL", f"Fatal error during execution: {e}")

finally:
    # --- Finalizing and cleaning up ---
    log_step(sheet, "Saving Excel log and closing browser", "INFO")
    if 'workbook' in locals() and 'driver' in locals():
        workbook.save(log_filename)
        driver.quit()
    elif 'workbook' in locals():
        workbook.save(log_filename)

    print(f"Automation log saved to {log_filename}")