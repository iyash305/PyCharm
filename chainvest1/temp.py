import time
from itertools import count
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment
from datetime import datetime

from selenium import webdriver
from selenium.webdriver.common.by import By
from selenium.webdriver.support.ui import Select
from selenium.webdriver.common.action_chains import ActionChains
from selenium.webdriver.support.ui import WebDriverWait
from selenium.webdriver.support import expected_conditions as EC


# A simple class for logging to an Excel file
class ExcelLogger:
    def __init__(self, filename="Chainvest_logs.xlsx"):
        self.filename = filename
        self.workbook = openpyxl.Workbook()
        self.sheet = self.workbook.active
        self.sheet.title = "Automation Log"
        self.log_row = 1
        self.setup_headers()

    def setup_headers(self):
        headers = ["Timestamp", "Action", "Log Message"]
        self.sheet.append(headers)

        # Apply formatting to headers
        header_font = Font(bold=True, color="FFFFFF")
        header_fill = PatternFill(start_color="4F81BD", end_color="4F81BD", fill_type="solid")

        for cell in self.sheet[1]:
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = Alignment(horizontal="center", vertical="center")

        # Set column widths for better readability
        self.sheet.column_dimensions['A'].width = 25
        self.sheet.column_dimensions['B'].width = 30
        self.sheet.column_dimensions['C'].width = 80
        self.log_row += 1

    def log(self, action, message, color=None):
        timestamp = datetime.now().strftime("%Y-%m-%d %H:%M:%S")
        self.sheet.append([timestamp, action, message])

        # Apply color coding if specified
        if color:
            fill = PatternFill(start_color=color, end_color=color, fill_type="solid")
            for cell in self.sheet[self.log_row]:
                cell.fill = fill

        self.log_row += 1

    def save(self):
        self.workbook.save(self.filename)


# Initialize the logger
logger = ExcelLogger(r"C:\Users\manoj\Desktop\Logs\Chainvest.xlsx")

# -----setting up browser----
options = webdriver.ChromeOptions()
options.add_argument("--incognito")
driver = webdriver.Chrome(options=options)
driver.maximize_window()
wait = WebDriverWait(driver, 20)

# ---navigating to login page and logging in---
logger.log("Navigation", "Navigating to login page.")
driver.get("http://code-server-app/Chainvest/BO/login")

logger.log("Login", "Entering credentials and logging in.")
email = wait.until(EC.presence_of_element_located((By.ID, "Email")))
email.send_keys("admin@chainvest.com")
password = driver.find_element(By.ID, "Password")
password.send_keys("P@ssw0rd")
sign_in = driver.find_element(By.XPATH, "//button[normalize-space(text())='Sign in']")
sign_in.click()

# ---Adding a new business
logger.log("Add Business", "Clicking Business menu.")
business_menu = wait.until(EC.element_to_be_clickable((By.XPATH, "//span[normalize-space(text())='Business']")))
business_menu.click()

logger.log("Add Business", "Clicking 'Add Business' button.")
add_busi = wait.until(EC.element_to_be_clickable((By.XPATH, "//button[contains(., 'Add Business')]")))
add_busi.click()

logger.log("Fill Business Details", "Filling in business information.")
wait.until(EC.presence_of_element_located((By.ID, "Name"))).send_keys("DMart")
business_type = driver.find_element(By.XPATH, "//select[@id='BusinessTypeID']")
select = Select(business_type)
select.select_by_visible_text("Retailer")
business_category = driver.find_element(By.XPATH, "//input[@role='combobox']")
business_category.send_keys("Retail")
option = wait.until(EC.element_to_be_clickable((By.XPATH, "//span[normalize-space()='Retail']")))
option.click()
media = driver.find_element(By.XPATH, "//input[@id='Logo']")
media.send_keys(r"C:\Users\manoj\Desktop\SampleImages\dmart.png")
driver.find_element(By.ID, "ContactPersonName").send_keys("DMart Admin")
driver.find_element(By.ID, "EmailID").send_keys("dmart@sample.com")
driver.find_element(By.ID, "MobileNo1").send_keys("99911212111")
driver.find_element(By.ID, "MobileNo2").send_keys("9786411087")
driver.find_element(By.ID, "WebsiteURL").send_keys("https://www.dmartindia.com/")
driver.find_element(By.ID, "BusinessLicenseNo").send_keys("A12345")
driver.find_element(By.ID, "VAT").send_keys("IE1234567T")
driver.find_element(By.ID, "Description").send_keys("Test description")
driver.find_element(By.XPATH, "//span[normalize-space()='Save & Next']").click()
time.sleep(0.5)

logger.log("Upload Document", "Uploading business document.")
media = driver.find_element(By.XPATH,"//input[@id='Logo']")

media.send_keys(r"C:\Users\manoj\Desktop\SampleImages\dmart.png")
doc_save = wait.until(EC.element_to_be_clickable((By.XPATH, "//button[span[normalize-space()='Save & Next']]")))
doc_save.click()
time.sleep(0.25)

logger.log("Address Details", "Entering address information.")
address1 = wait.until(
    EC.presence_of_element_located((By.XPATH, "//input[@type='text' and @formcontrolname='Address1']")))
address1.send_keys("Bangur Nagar")
address2 = driver.find_element(By.XPATH, "//input[@type='text' and @formcontrolname='Address2']")
address2.send_keys("Goregaon West")

country = driver.find_element(By.XPATH, "//select[@formcontrolname='CountryID']")
select_country = Select(country)
select_country.select_by_visible_text("Ethiopia")
logger.log("Address Details", "Selected Country: Ethiopia.")

wait.until(EC.visibility_of_element_located(
    (By.XPATH, "//select[@formcontrolname='StateID']/option[normalize-space()='Addis Ababa']")))
state = wait.until(EC.element_to_be_clickable((By.XPATH, "//select[@formcontrolname='StateID']")))
select_state = Select(state)
select_state.select_by_visible_text("Addis Ababa")
logger.log("Address Details", "Selected State: Addis Ababa.")

city_dropdown = wait.until(EC.element_to_be_clickable((By.XPATH, "//select[@formcontrolname='CityID']")))
select_city = Select(city_dropdown)
wait.until(EC.visibility_of_element_located(
    (By.XPATH, "//select[@formcontrolname='CityID']/option[normalize-space()='Addis Ababa']")))
select_city.select_by_visible_text("Addis Ababa")
logger.log("Address Details", "Selected City: Addis Ababa.")

driver.find_element(By.XPATH, "//input[@formcontrolname='Zip']").send_keys("400123")
driver.find_element(By.XPATH, "//input[@formcontrolname='Landmark']").send_keys("Addis")
driver.find_element(By.XPATH, "//input[@formcontrolname='Latitude']").send_keys("40.12")
driver.find_element(By.XPATH, "//input[@formcontrolname='Longitude']").send_keys("35.98")
driver.find_element(By.XPATH, "//span[normalize-space()='Save & Next']").click()
time.sleep(0.5)

logger.log("Bank Details", "Entering bank information.")
bank_name = wait.until(EC.presence_of_element_located((By.XPATH, "//input[@formcontrolname='BankName']")))
bank_name.send_keys("HDFC")
driver.find_element(By.XPATH, "//input[@formcontrolname='AccountHolderName']").send_keys("DMart")
driver.find_element(By.XPATH, "//input[@formcontrolname= 'AccountNumber']").send_keys("2874271980167")
driver.find_element(By.XPATH, "//input[@formcontrolname= 'IFSC']").send_keys("HDFC00000504")
driver.find_element(By.XPATH, "//input[@formcontrolname= 'BankBranchName']").send_keys("HDFC Addis")
finish = driver.find_element(By.XPATH, "//span[normalize-space()='Save & Complete']")
finish.click()
logger.log("Business", "Business details saved successfully!", color="C6EFCE")  # Green for success
time.sleep(1)

# ---Master Menu
logger.log("Navigation", "Navigating to Master menu.")
wait.until(EC.presence_of_element_located((By.XPATH, "//span[normalize-space() = 'Master']"))).click()
wait.until(EC.presence_of_element_located((By.XPATH, "//a[@href='/Chainvest/BO/business-type']"))).click()
add_bus_type = wait.until(
    EC.presence_of_element_located((By.XPATH, "//button[normalize-space()= 'Add Business Type']"))).click()
busi_type = driver.find_element(By.XPATH, "//input[@formcontrolname='BusinessTypeName']")
busi_type.send_keys("Agent")
busi_desc = driver.find_element(By.XPATH, "//textarea[@formcontrolname='Description']")
busi_desc.send_keys("Acts on behalf of a company to sell or promote products/services, usually earning a commission.")
driver.find_element(By.XPATH, "//button[normalize-space(text())='Add']").click()
logger.log("Master", "Business Type 'Agent' added successfully.", color="C6EFCE")

# ---Language
#logger.log("Navigation", "Navigating to Languages.")
#driver.find_element(By.XPATH, "//a[@href='/Chainvest/BO/language']").click()
#add_lang_button = driver.find_element(By.XPATH, "//button[normalize-space(text())='Add Language']").click()
#ent_lang = driver.find_element(By.ID, "name")
#ent_lang.send_keys("French")
#lang_desc = driver.find_element(By.XPATH, "//textarea[@id='description']")
#lang_desc.send_keys("Français")
#lang_code = driver.find_element(By.XPATH, "//input[@formcontrolname='Languagecode']")
#lang_code.send_keys("FR")
#lc_id = driver.find_element(By.XPATH, "//input[@formcontrolname='LCID']")
#lc_id.send_keys("33")
#add_lang = driver.find_element(By.XPATH, "//button[normalize-space(text())='Add']").click()
#logger.log("Master", "Language 'French' added successfully.", color="C6EFCE")

# ---Roles---
logger.log("Navigation", "Navigating to Roles.")
driver.find_element(By.XPATH, "//a[@href='/Chainvest/BO/roles']").click()
role_add = driver.find_element(By.XPATH, "//button[normalize-space(text())='Add Role']").click()
select_role = driver.find_element(By.XPATH, "//select[@formcontrolname='PrimaryRoleID']")
select = Select(select_role)
select.select_by_visible_text("EndUser")
role1 = driver.find_element(By.XPATH, "//input[@formcontrolname='Name']")
role1.send_keys("End User")
role_desc = driver.find_element(By.XPATH, "//textarea[@formcontrolname='Description']")
role_desc.send_keys("Application User")
add_role = driver.find_element(By.XPATH, "//button[normalize-space(text())='Add']").click()
logger.log("Master", "Role 'End User' added successfully.", color="C6EFCE")

# ---Payment Type
logger.log("Navigation", "Navigating to Payment Types.")
driver.find_element(By.XPATH, "//a[@href='/Chainvest/BO/paymentType']").click()
add_pymnt_type = driver.find_element(By.XPATH, "//button[normalize-space(text())='Add Payment Type']").click()
pymnt_type = driver.find_element(By.XPATH, "//input[@formcontrolname='PaymentTypeName']")
pymnt_type.send_keys("QR Code")
pymnt_desc = driver.find_element(By.XPATH, "//textarea[@formcontrolname='PaymentTypeDescription']")
pymnt_desc.send_keys("Scan an QR code for faster payment")
add_pymnt = driver.find_element(By.XPATH, "//button[normalize-space(text())='Add']").click()
logger.log("Master", "Payment Type 'QR Code' added successfully.", color="C6EFCE")

# ---Payment Status
logger.log("Navigation", "Navigating to Payment Status.")
driver.find_element(By.XPATH, "//a[@href='/Chainvest/BO/paymentStatus']").click()
payment_status_button = driver.find_element(By.XPATH, "//button[normalize-space(text())='Add Payment Status']").click()
enter_payment_status = driver.find_element(By.XPATH, "//input[@formcontrolname='StatusName']")
enter_payment_status.send_keys("Initiated")
add_pymnt_button = driver.find_element(By.XPATH, "//button[normalize-space(text())='Add']").click()
logger.log("Master", "Payment Status 'Initiated' added successfully.", color="C6EFCE")

# ---Order Status---#
logger.log("Navigation", "Navigating to Order Status.")
driver.find_element(By.XPATH, "//a[@href='/Chainvest/BO/order-status']").click()
add_order_status = driver.find_element(By.XPATH, "//button[normalize-space(text())='Add Order Status']")
add_order_status.click()
enter_order_status = driver.find_element(By.ID, "OrderStatusName")
enter_order_status.send_keys("Pending")
enter_order_desc = driver.find_element(By.ID, "OrderStatusDescription")
enter_order_desc.send_keys("Order has been placed but not yet processed.")
add_status = driver.find_element(By.XPATH, "//button[normalize-space(text())='Add']").click()
logger.log("Master", "Order Status 'Pending' added successfully.", color="C6EFCE")

# ---Loan Setting---#
logger.log("Navigation", "Navigating to Loan Settings.")
loan_module = wait.until(EC.presence_of_element_located((By.XPATH, "//span[normalize-space(text())='Loan Setting']")))
loan_module.click()

loan_purpose_module = driver.find_element(By.XPATH, "//a[@href='/Chainvest/BO/loan-purpose']").click()
add_loan_purpose_btn = driver.find_element(By.XPATH, "//button[normalize-space(text())='Add Loan Purpose']")
add_loan_purpose_btn.click()
enter_loan_purpose = driver.find_element(By.ID, "LoanPurposeName")
enter_loan_purpose.send_keys("Logistics")
enter_loan_desc = driver.find_element(By.ID, "LoanPurposeDescription")
enter_loan_desc.send_keys("Fleet expansion or delivery service payments")
int_from = driver.find_element(By.XPATH, "//input[@formcontrolname='InterestFrom']")
int_from.send_keys("8")
int_to = driver.find_element(By.XPATH, "//input[@formcontrolname='InterestTO']")
int_to.send_keys("20")
loan_amount_from = driver.find_element(By.ID, "name")
loan_amount_from.send_keys("200000")
loan_amount_to = driver.find_element(By.ID, "LoanAmountTO")
loan_amount_to.send_keys("800000")
add_loan_purpose = driver.find_element(By.XPATH, "//button[normalize-space(text())='Add']")
add_loan_purpose.click()
logger.log("Loan Setting", "Loan Purpose 'Logistics' added successfully.", color="C6EFCE")

loan_status_module = driver.find_element(By.XPATH, "//a[@href='/Chainvest/BO/loan-status']").click()
add_loan_status_btn = driver.find_element(By.XPATH, "//button[normalize-space(text())='Add Loan Status']")
add_loan_status_btn.click()
add_loan_status = driver.find_element(By.ID, "LoanStatusName")
add_loan_status.send_keys("Overdue")
loan_status_desc = driver.find_element(By.ID, "LoanStatusDescription")
loan_status_desc.send_keys("Payment is delayed beyond the due date.")
add_btn = driver.find_element(By.XPATH, "//button[normalize-space(text())='Add']")
add_btn.click()
logger.log("Loan Setting", "Loan Status 'Overdue' added successfully.", color="C6EFCE")

# ---Credit Score---#
logger.log("Navigation", "Navigating to Credit Score.")
credit_score = driver.find_element(By.XPATH, "//span[normalize-space(text())='Credit Score']")
credit_score.click()
view_rule = driver.find_element(By.XPATH, "//a[@href='/Chainvest/BO/credit-score']").click()

scrollable_module = driver.find_element(By.CLASS_NAME, "overflow-y-scroll")
driver.execute_script("arguments[0].scrollTop = arguments[0].scrollHeight;", scrollable_module)
time.sleep(0.5)
driver.execute_script("arguments[0].scrollTop = 0;", scrollable_module)

add_rule_button = driver.find_element(By.XPATH, "//button[text()=' Add Rule ']")
add_rule_button.click()
time.sleep(2)
rule_name = driver.find_element(By.XPATH, "//input[@id='RuleName']")
rule_name.send_keys("Credit Utilization")
rule_desc = driver.find_element(By.XPATH, "//textarea[@id='Description']")
rule_desc.send_keys("Credit utilization ratio between 10% and 30%")
score = driver.find_element(By.ID, "ScoreImpact")
score.send_keys("10")
prio = driver.find_element(By.ID, "Priority")
prio.clear()
prio.send_keys("4")
add_rule = driver.find_element(By.XPATH, "//button[normalize-space(text())='Create Rule']")
add_rule.click()
logger.log("Credit Score", "Credit Rule 'Credit Utilization' added successfully.", color="C6EFCE")

view_score = wait.until(
    EC.presence_of_element_located((By.XPATH, "//a[@href='/Chainvest/BO/credit-scorescreen']"))).click()
view_score_company = wait.until(
    EC.presence_of_element_located((By.XPATH, "//li[normalize-space()='InnovatePro Fabrications']")))
view_score_company.click()
logger.log("Credit Score", "Viewing Credit Score for InnovatePro Fabrications.", color="C6EFCE")

# ---Loan---#
logger.log("Navigation", "Navigating to Loans.")
loan = driver.find_element(By.XPATH, "//span[normalize-space(text())='Loan']")
loan.click()
view = wait.until(EC.element_to_be_clickable((By.XPATH, "//button[normalize-space()='View']")))
view.click()
driver.execute_script("window.scrollTo(0, document.body.scrollHeight);")
time.sleep(1)
driver.execute_script("window.scrollTo(0, 0);")
time.sleep(2)
logger.log("Loan", "Viewed loan details.", color="C6EFCE")

# ---Product---#
logger.log("Navigation", "Navigating to Product menu.")
product = driver.find_element(By.XPATH, "//span[normalize-space()='Product']")
product.click()

bulk_upload_click = driver.find_element(By.XPATH, "//a[@href='/Chainvest/BO/bulk-upload']").click()
file_input = WebDriverWait(driver, 10).until(EC.presence_of_element_located((By.ID, "fileInput")))
file_input.send_keys(r"C:\Users\manoj\Downloads\chainvestupload.xlsx")
fileupload = driver.find_element(By.XPATH, "//button[normalize-space()='Upload File']")
fileupload.click()
logger.log("Product", "Bulk upload file submitted.", color="C6EFCE")

product_category = driver.find_element(By.XPATH, "//a[@href='/Chainvest/BO/product-category']").click()
add_product_category_btn = wait.until(
    EC.element_to_be_clickable((By.XPATH, "//button[normalize-space(text())='Add Product Category']")))
add_product_category_btn.click()

logger.log("Product", "Adding new product category: Earbuds.")
thumb_img = wait.until(EC.presence_of_element_located((By.XPATH, "(//input[@type='file'])[1]")))
thumb_img.send_keys(r"C:\Users\manoj\Downloads\earbuds.jpeg")
large_input = driver.find_element(By.XPATH, "(//input[@type='file'])[2]")
large_input.send_keys(r"C:\Users\manoj\Downloads\earbuds.jpeg")

# Fill in details for different languages
languages = [
    ("Hindi", "Earbuds", "Earbuds"),
    ("Afan Oromo", "Earbuds", "Earbuds"),
    ("Amharic", "Earbuds", "Earbuds"),
    ("English", "Earbuds", "Earbuds")
]

for lang_name, short_name, desc in languages:
    name_xpath = f"//tr[td[normalize-space()='{lang_name}']]//input[@formcontrolname='ProductCategoryShortName']"
    desc_xpath = f"//tr[td[normalize-space()='{lang_name}']]//input[@placeholder='Enter Description']"

    try:
        name_field = wait.until(EC.presence_of_element_located((By.XPATH, name_xpath)))
        name_field.send_keys(short_name)
        desc_field = driver.find_element(By.XPATH, desc_xpath)
        desc_field.send_keys(desc)
        logger.log("Product Category", f"Filled details for '{lang_name}'.", color="C6EFCE")
    except Exception as e:
        logger.log("Error", f"Could not fill details for '{lang_name}': {str(e)}", color="FFC7CE")

add_btn = driver.find_element(By.XPATH, "//button[normalize-space()='Add']")
add_btn.click()
logger.log("Product", "Product Category added successfully.", color="C6EFCE")

# ---Product SubCategory---#
sub_category = wait.until(EC.element_to_be_clickable((By.XPATH, "//a[@href='/Chainvest/BO/product-subcategory']")))
sub_category.click()

sub_category_btn = wait.until(EC.element_to_be_clickable((By.XPATH, "//button[text()='Add Product SubCategory']")))
sub_category_btn.click()

sub_category_dropdown = wait.until(
    EC.presence_of_element_located((By.XPATH, "//select[@formcontrolname='ProductCategoryID']")))
select_sub_category = Select(sub_category_dropdown)
select_sub_category.select_by_visible_text("Earbuds")
logger.log("Product SubCategory", "Selected 'Earbuds' as product category.", color="C6EFCE")

sub_category_upload = driver.find_element(By.XPATH, "(//input[@type='file'])[1]")
sub_category_upload.send_keys(r"C:\Users\manoj\Downloads\earbuds.jpeg")
sub_category_upload1 = driver.find_element(By.XPATH, "(//input[@type='file'])[2]")
sub_category_upload1.send_keys(r"C:\Users\manoj\Downloads\earbuds.jpeg")

# Fill in sub-category details for different languages
sub_cat_languages = [
    ("Hindi", "Wireless earbuds", "Wireless earbuds"),
    ("Afan Oromo", "Wireless earbuds", "Wireless earbuds"),
    ("Amharic", "Wireless earbuds", "Wireless earbuds"),
    ("English", "Wireless Earbuds", "Wireless earbuds")
]

for lang_name, short_name, desc in sub_cat_languages:
    name_xpath = f"//tr[td[normalize-space()='{lang_name}']]//input[@formcontrolname='ProductSubCategoryShortName']"
    desc_xpath = f"//tr[td[normalize-space()='{lang_name}']]//input[@placeholder='Enter Description']"

    try:
        name_field = wait.until(EC.presence_of_element_located((By.XPATH, name_xpath)))
        name_field.send_keys(short_name)
        desc_field = driver.find_element(By.XPATH, desc_xpath)
        desc_field.send_keys(desc)
        logger.log("Product SubCategory", f"Filled details for '{lang_name}'.", color="C6EFCE")
    except Exception as e:
        logger.log("Error", f"Could not fill details for '{lang_name}' in sub-category: {str(e)}", color="FFC7CE")

add_sub_cat = driver.find_element(By.XPATH, "//button[normalize-space()='Add']")
add_sub_cat.click()
logger.log("Product SubCategory", "Product SubCategory added successfully.", color="C6EFCE")

# ---Orders---#
orders = wait.until(EC.presence_of_element_located((By.XPATH, "//span[normalize-space(text())='Orders']")))
orders.click()
view = wait.until(EC.presence_of_element_located((By.XPATH, "//button[normalize-space(text())='View']")))
view.click()
logger.log("Orders", "Viewing order details.", color="C6EFCE")

# ---Users---#
user_module = wait.until(EC.presence_of_element_located((By.XPATH, "//span[normalize-space(text())='Users']")))
user_module.click()
new_user_add_btn = wait.until(
    EC.presence_of_element_located((By.XPATH, "//button[normalize-space(text())='Add User']")))
new_user_add_btn.click()
logger.log("Users", "Adding new user.", color="C6EFCE")

company_dropdown = wait.until(EC.presence_of_element_located((By.XPATH, "//select[@formcontrolname = 'CompanyID']")))
select_company = Select(company_dropdown)
select_company.select_by_visible_text("Digital Automation Enterprises")
driver.find_element(By.XPATH, "//input[@formcontrolname = 'Email']").send_keys("admin@dae.com")
driver.find_element(By.XPATH, "//input[@formcontrolname = 'PasswordHash']").send_keys("P@ssw0rd")
user_country_dropdown = driver.find_element(By.XPATH, "//select[@formcontrolname = 'CountryID']")
country_select = Select(user_country_dropdown)
country_select.select_by_visible_text("Ethiopia")
driver.find_element(By.XPATH, "//input[@formcontrolname = 'Phone']").send_keys("9874615860")
role_dropdown = driver.find_element(By.XPATH, "//select[@formcontrolname = 'RoleID']")
role_select = Select(role_dropdown)
role_select.select_by_visible_text("Admin")
add_user = driver.find_element(By.XPATH, "//button[normalize-space(text())='Add']")
add_user.click()
logger.log("Users", "New user added successfully.", color="C6EFCE")

# ---Reports---#
reports_module = wait.until(EC.presence_of_element_located((By.XPATH, "//span[normalize-space(text())='Reports']")))
reports_module.click()
order_report = wait.until(EC.presence_of_element_located((By.XPATH, "//a[@href='/Chainvest/BO/fulfillment-report']")))
order_report.click()
date_from = wait.until(EC.presence_of_element_located((By.XPATH, "//input[@formcontrolname='orderDateFrom']")))
date_from.send_keys("01-08-2025")
date_to = wait.until(EC.presence_of_element_located((By.XPATH, "//input[@formcontrolname='orderDateTo']")))
date_to.send_keys("20-08-2025")
search = driver.find_element(By.XPATH, "//button[normalize-space(text())='Search']")
search.click()
time.sleep(2)
clear_btn = driver.find_element(By.XPATH, "//button[normalize-space(text())='Clear']")
clear_btn.click()
logger.log("Reports", "Order report search performed and cleared.", color="C6EFCE")

bank_reports = driver.find_element(By.XPATH, "//a[@href='/Chainvest/BO/bankandsupplychain-report']")
bank_reports.click()
time.sleep(2)
financial_reports = driver.find_element(By.XPATH, "//a[@href='/Chainvest/BO/financial-report']")
financial_reports.click()
time.sleep(3)
logger.log("Reports", "Financial reports viewed.", color="C6EFCE")

# Save the log file before closing the browser
logger.save()

# Close the browser at the end
driver.quit()